#!/usr/bin/env python3
"""Build every other form of the résumé from resume/index.html, word for word.

resume/index.html is the only place the words live. Its Designed / Plain switch restyles those
same words in place, and this tool copies them into five files:

    resume/plain/index.html               the plain text page
    resume/Mike_Purpura_Resume.pdf        the designed view, printed (two letter pages)
    resume/Mike_Purpura_Resume_Plain.pdf  the plain view, printed (two letter pages)
    resume/Mike_Purpura_Resume.docx       Word, laid out like the designed view
    resume/Mike_Purpura_Resume_Plain.docx Word, plain: one column, real headings and lists

Then it reads every file back and stops if a single word differs from resume/index.html or a PDF
isn't exactly two pages. It stops on markup it doesn't know instead of guessing, so a new kind of
block in the résumé needs a line here too. Edit resume/index.html, then rerun:

    uv run --with python-docx --with pymupdf --with websocket-client python3 tools/resume_build.py
    ... tools/resume_build.py --check    # change nothing; exit 1 if any file is out of date

The PDFs are printed by headless Brave (or Chrome) from a throwaway local server. Only the QR code
and the diamond bullets stay out of the plain forms, because they're pictures, not words; the two
lines at the foot of the page are screen-only, so the printed forms leave them out too.
"""
import argparse
import base64
import difflib
import functools
import html
import http.server
import io
import json
import os
import re
import shutil
import struct
import subprocess
import sys
import tempfile
import threading
import time
import urllib.request
import zlib
from datetime import datetime, timezone
from html.parser import HTMLParser
from urllib.parse import urljoin

SITE = "https://okay-cool.net"
DESIGNED_URL = SITE + "/resume/"
REPO = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
SOURCE = "resume/index.html"
PLAIN_PAGE = "resume/plain/index.html"
DESIGNED_PDF = "resume/Mike_Purpura_Resume.pdf"
PLAIN_PDF = "resume/Mike_Purpura_Resume_Plain.pdf"
DESIGNED_DOCX = "resume/Mike_Purpura_Resume.docx"
PLAIN_DOCX = "resume/Mike_Purpura_Resume_Plain.docx"
QR = "resume/qr-okay-cool.png"
PDF_PAGES = 2
BROWSERS = ["/Applications/Brave Browser.app/Contents/MacOS/Brave Browser",
            "/Applications/Google Chrome.app/Contents/MacOS/Google Chrome",
            "chromium", "google-chrome", "brave-browser"]

VOID = {"area", "base", "br", "col", "embed", "hr", "img", "input", "link", "meta", "source", "track", "wbr"}
BLOCK = {"address", "article", "br", "dd", "div", "dl", "dt", "figcaption", "figure", "footer", "h1", "h2", "h3",
         "h4", "header", "li", "main", "nav", "ol", "p", "section", "ul"}
DOWNLOADS = [("Designed PDF", DESIGNED_PDF), ("Designed DOCX", DESIGNED_DOCX),
             ("Plain PDF", PLAIN_PDF), ("Plain DOCX", PLAIN_DOCX)]


# ------------------------------------------------------------------ reading resume/index.html

class Node(object):
    def __init__(self, tag, attrs, parent):
        self.tag, self.attrs, self.parent, self.children = tag, dict(attrs), parent, []

    @property
    def classes(self):
        return set((self.attrs.get("class") or "").split())

    def is_(self, tag, cls=None):
        return self.tag == tag and (cls is None or cls in self.classes)

    def elements(self):
        return [c for c in self.children if isinstance(c, Node)]

    def __repr__(self):
        cls = self.attrs.get("class")
        return "<%s%s>" % (self.tag, ' class="%s"' % cls if cls else "")


class Tree(HTMLParser):
    def __init__(self, text):
        HTMLParser.__init__(self, convert_charrefs=True)
        self.root = self.cur = Node("#root", [], None)
        self.feed(text)
        self.close()

    def handle_starttag(self, tag, attrs):
        node = Node(tag, attrs, self.cur)
        self.cur.children.append(node)
        if tag not in VOID:
            self.cur = node

    def handle_startendtag(self, tag, attrs):
        self.cur.children.append(Node(tag, attrs, self.cur))

    def handle_endtag(self, tag):
        if tag in VOID:
            return
        node = self.cur
        while node.parent is not None and node.tag != tag:
            node = node.parent
        if node.parent is None:
            raise ValueError("unmatched </%s>" % tag)
        self.cur = node.parent

    def handle_data(self, data):
        self.cur.children.append(data)


class Unknown(ValueError):
    def __init__(self, node):
        ValueError.__init__(self, "tools/resume_build.py doesn't know how to copy %r (inside %r); teach it, "
                                  "don't guess" % (node, node.parent))


def find(node, pred):
    for c in node.elements():
        if pred(c):
            yield c
        for x in find(c, pred):
            yield x


def the_one(node, pred, what):
    found = list(find(node, pred))
    if len(found) != 1:
        raise ValueError("expected one %s in the résumé, found %d" % (what, len(found)))
    return found[0]


def is_glyph(node):
    """An <i> holding only a symbol (the ◆ before each contact item) is decoration, not words."""
    text = "".join(c for c in node.children if isinstance(c, str)).strip()
    return node.tag == "i" and not node.elements() and bool(text) and not any(ch.isalnum() for ch in text)


def left_out(node):
    return node.is_("figure", "qr") or is_glyph(node) or node.tag in ("script", "style", "svg", "template")


def words(node, printed=False, email=None):
    """The words under a node, in order. Blocks split words, and so do tags that sit side by side
    with no text between them (the designed page lays those out as separate flex items); a tag
    inside running text doesn't. printed=True gives what a printout shows: the email address
    instead of its placeholder, no <noscript>, nothing marked no-print."""
    out = []

    def walk(n):
        apart = not any(isinstance(c, str) and c.strip() for c in n.children)
        for c in n.children:
            if isinstance(c, str):
                out.append(c)
            elif left_out(c) or (printed and (c.tag == "noscript" or "no-print" in c.classes)):
                continue
            else:
                gap = " " if c.tag in BLOCK or apart else ""
                out.append(gap)
                if printed and c.attrs.get("id") == "em":
                    out.append(email)
                else:
                    walk(c)
                out.append(gap)

    walk(node)
    return " ".join("".join(out).split())


def only_tags(node):
    for c in node.children:
        if isinstance(c, str) and c.strip():
            raise ValueError("loose text %r directly inside %r" % (c.strip()[:40], node))
    return node.elements()


def read_source(text):
    src = Tree(text).root
    sheet = the_one(src, lambda n: n.is_("main", "sheet"), '<main class="sheet">')
    hidden = the_one(sheet, lambda n: n.tag == "noscript", "<noscript> email")
    spelled = words(hidden)
    email = re.sub(r"\s*\[dot\]\s*", ".", re.sub(r"\s*\[at\]\s*", "@", spelled))
    if not re.fullmatch(r"[\w.+-]+@[\w-]+(\.[\w-]+)+", email):
        raise ValueError("can't read the email address out of %r" % spelled)
    return src, sheet, email


# ------------------------------------------------------------------ the plain text page

CSS = """\
:root { color-scheme: light dark; --bg: #fff; --ink: #111; --dim: #3b3b3b; --line: #c9c9c9; --link: #0b4fb3; }
@media (prefers-color-scheme: dark) { :root { --bg: #121212; --ink: #f2f2f2; --dim: #d0d0d0; --line: #4d4d4d; --link: #9cc3ff; } }
* { box-sizing: border-box; }
html { -webkit-text-size-adjust: 100%; }
body { margin: 0; background: var(--bg); color: var(--ink); font: 18px/1.6 system-ui, -apple-system, "Segoe UI", Roboto, Helvetica, Arial, sans-serif; }
nav, main { max-width: 46rem; margin: 0 auto; padding-inline: 20px; }
nav { display: flex; flex-wrap: wrap; align-items: center; gap: 8px 22px; font-size: 16px; }
.top { padding-block: 14px; border-bottom: 1px solid var(--line); }
.top button { margin-left: auto; font: inherit; color: var(--ink); background: var(--bg); border: 1px solid currentColor; border-radius: 6px; padding: 3px 14px; cursor: pointer; }
@media (max-width: 520px) { .top button { margin-left: 0; } }
.downloads { gap: 8px 20px; padding-bottom: 48px; }
.dl-k { font-weight: 700; }
a { color: var(--link); text-underline-offset: 2px; overflow-wrap: anywhere; }
:focus-visible { outline: 3px solid var(--link); outline-offset: 2px; }
main { padding-block: 28px 24px; }
h1 { font-size: 2rem; line-height: 1.15; margin: 0 0 6px; }
h2 { font-size: 1.3rem; line-height: 1.3; margin: 34px 0 12px; padding-bottom: 4px; border-bottom: 2px solid var(--ink); }
h3 { font-size: 1.1rem; line-height: 1.35; margin: 24px 0 2px; }
h2 + h3, h2 + .role > h3:first-child { margin-top: 0; }
p, ul, dl { margin: 0 0 12px; }
ul { padding-left: 1.3em; }
li { margin-bottom: 6px; }
.title { font-weight: 600; }
.contact { list-style: none; padding: 0; display: flex; flex-wrap: wrap; gap: 2px 24px; }
.contact li { margin: 0; }
.meta { margin: 0; color: var(--dim); }
.meta + ul, .meta + p:not(.meta) { margin-top: 8px; }
dt { font-weight: 700; }
dd { margin: 0 0 12px; }
footer { margin-top: 40px; padding-top: 12px; border-top: 1px solid var(--line); color: var(--dim); }
footer p { margin: 0 0 4px; }
@media print {
  @page { size: letter; margin: 0.5in 0.6in; }
  :root { --bg: #fff; --ink: #000; --dim: #000; --line: #888; --link: #000; }
  body { font-size: 10.5pt; line-height: 1.4; }
  .no-print { display: none !important; }
  main { max-width: none; padding: 0; }
  h1 { font-size: 20pt; }
  h2 { font-size: 12.5pt; margin: 14pt 0 6pt; break-after: avoid; }
  h3 { font-size: 11pt; margin: 10pt 0 1pt; break-after: avoid; }
  .meta { break-after: avoid; }
  li { margin-bottom: 3pt; break-inside: avoid; }
  a { text-decoration: none; }
}"""

PAGE = """\
<!doctype html>
<html lang="en">
<head>
<meta charset="utf-8">
<meta name="viewport" content="width=device-width, initial-scale=1">
<title>@TITLE@</title>
<meta name="description" content="@DESCRIPTION@">
<meta name="color-scheme" content="light dark">
<link rel="canonical" href="https://okay-cool.net/resume/">
<link rel="icon" href="/favicon.svg" type="image/svg+xml">
<!-- The plain text version of the résumé, word for word from ../index.html (where the words live).
     Built by tools/resume_build.py: edit ../index.html, then rerun it. Don't edit this file by hand. -->
<style>
@CSS@
</style>
</head>
<body>
<nav class="top no-print" aria-label="Résumé">
  <a href="/">Back to okay-cool.net</a>
  <a href="/resume/">Designed résumé</a>
  <button type="button" id="print-btn">Print</button>
</nav>
<main>
@MAIN@
</main>
<nav class="downloads no-print" aria-label="Download the résumé">
  <span class="dl-k">Download</span>
@DOWNLOADS@
</nav>
<script src="/assets/js/resume.js" defer></script>
</body>
</html>
"""


def esc(s):
    return html.escape(s, quote=False)


def link(value):
    """The designed page lives at /resume/; make its relative links work from /resume/plain/ too."""
    if value.startswith("#"):
        return value
    url = urljoin(DESIGNED_URL, value)
    return url[len(SITE):] if url.startswith(SITE + "/") else value


def inline(node):
    out = []
    for c in node.children:
        if isinstance(c, str):
            out.append(esc(c))
        elif is_glyph(c):
            continue
        elif c.tag == "span":
            out.append(inline(c))
        elif c.tag in ("strong", "em", "b"):
            out.append("<%s>%s</%s>" % (c.tag, inline(c), c.tag))
        elif c.tag == "a":
            ident = ' id="%s"' % html.escape(c.attrs["id"]) if c.attrs.get("id") else ""
            out.append('<a%s href="%s">%s</a>' % (ident, html.escape(link(c.attrs.get("href") or "")), inline(c)))
        elif c.tag == "br":
            out.append("<br>")
        elif c.tag == "noscript":
            out.append("<noscript>%s</noscript>" % inline(c))
        else:
            raise Unknown(c)
    return "".join(out)


def slug(text):
    return re.sub(r"[^a-z0-9]+", "-", text.lower()).strip("-")


def _header(node):
    out = ["<header>"]

    def walk(n):
        for c in only_tags(n):
            if c.tag == "div" and not c.classes:
                walk(c)
            elif c.is_("figure", "qr"):
                continue  # the QR code is a picture of the portfolio URL, which the contact line spells out
            elif c.tag == "h1":
                out.append("<h1>%s</h1>" % inline(c))
            elif c.is_("p", "title"):
                out.append('<p class="title">%s</p>' % inline(c))
            elif c.is_("p", "contact"):
                out.append('<ul class="contact">')
                for item in only_tags(c):
                    if item.tag != "span":
                        raise Unknown(item)
                    out.append("<li>%s</li>" % inline(item).strip())
                out.append("</ul>")
            else:
                raise Unknown(c)

    walk(node)
    out.append("</header>")
    return out


def _bullets(node):
    items = []
    for li in only_tags(node):
        if li.tag != "li":
            raise Unknown(li)
        items.append("<li>%s</li>" % inline(li))
    return "<ul>\n%s\n</ul>" % "\n".join(items)


def _terms(node):
    rows = []
    for c in only_tags(node):
        if c.tag not in ("dt", "dd"):
            raise Unknown(c)
        rows.append("<%s>%s</%s>" % (c.tag, inline(c), c.tag))
    return "<dl>\n%s\n</dl>" % "\n".join(rows)


def _role(node):
    out = ['<article class="role">']
    for c in only_tags(node):
        if c.is_("div", "role-head"):
            for h in only_tags(c):
                if h.tag == "h3":
                    out.append("<h3>%s</h3>" % inline(h))
                elif h.tag == "span":
                    out.append('<p class="meta">%s</p>' % inline(h))
                else:
                    raise Unknown(h)
        elif c.is_("p", "org"):
            out.append('<p class="meta">%s</p>' % inline(c))
        elif c.tag == "ul":
            out.append(_bullets(c))
        else:
            raise Unknown(c)
    out.append("</article>")
    return "\n".join(out)


def _school(node):
    out = []
    for c in only_tags(node):
        if c.tag == "span" and "edu-when" in c.classes:
            out.append('<p class="meta">%s</p>' % inline(c))
        elif c.tag == "span" and not c.classes:
            out.append("<h3>%s</h3>" % inline(c))
        else:
            raise Unknown(c)
    return out


def _foot(node):
    lines = []
    for c in only_tags(node):
        if c.tag != "span":
            raise Unknown(c)
        lines.append("<p>%s</p>" % inline(c))
    cls = ' class="no-print"' if "no-print" in node.classes else ""
    return "<footer%s>\n%s\n</footer>" % (cls, "\n".join(lines))


def _plain_body(sheet):
    out, open_section = [], [False]

    def close():
        if open_section[0]:
            out.append("</section>")
            open_section[0] = False

    for n in only_tags(sheet):
        if n.is_("header", "top"):
            out.extend(_header(n))
        elif n.tag == "h2":
            close()
            sid = slug(words(n))
            out.append('<section aria-labelledby="%s">' % sid)
            out.append('<h2 id="%s">%s</h2>' % (sid, inline(n)))
            open_section[0] = True
        elif n.tag == "p" and n.classes <= {"summary", "edu-lines"}:
            out.append("<p>%s</p>" % inline(n))
        elif n.is_("dl", "skills"):
            out.append(_terms(n))
        elif n.is_("section", "role"):
            out.append(_role(n))
        elif n.is_("div", "edu"):
            out.extend(_school(n))
        elif n.tag == "ul":
            out.append(_bullets(n))
        elif n.is_("div", "foot"):
            close()
            out.append(_foot(n))
        else:
            raise Unknown(n)
    close()
    return out


def plain_page(source_html):
    """The plain text page for this résumé. Raises if any word would change."""
    src, sheet, _ = read_source(source_html)
    title = words(the_one(src, lambda n: n.tag == "title", "<title>"))
    desc = the_one(src, lambda n: n.is_("meta") and n.attrs.get("name") == "description", "meta description")
    downloads = "\n".join('  <a href="/%s" download>%s</a>' % (path, label) for label, path in DOWNLOADS)
    page = (PAGE.replace("@TITLE@", esc(title + " · Plain text"))
                .replace("@DESCRIPTION@", html.escape(desc.attrs.get("content") or ""))
                .replace("@CSS@", CSS)
                .replace("@DOWNLOADS@", downloads)
                .replace("@MAIN@", "\n".join(_plain_body(sheet))))
    same_words("the plain text page", words(sheet),
               words(the_one(Tree(page).root, lambda n: n.tag == "main", "<main> in the plain page")))
    return page


def same_words(what, want, got):
    if want != got:
        diff = "\n".join(difflib.unified_diff(want.split(), got.split(), "resume/index.html", what, lineterm="", n=4))
        raise ValueError("%s doesn't match resume/index.html word for word:\n%s" % (what, diff))


# ------------------------------------------------------------------ the words as a printout has them

class Run(object):
    """A piece of running text: plain, a link, the title's accent, or a line break."""
    def __init__(self, text="", href=None, accent=False, br=False):
        self.text, self.href, self.accent, self.br = text, href, accent, br


def _runs(node, email, accent=False):
    out = []
    for c in node.children:
        if isinstance(c, str):
            out.append(Run(c, accent=accent))
        elif is_glyph(c) or c.tag == "noscript":
            continue
        elif c.tag == "span":
            out.extend(_runs(c, email, accent=True))
        elif c.tag in ("strong", "em", "b"):
            out.extend(_runs(c, email, accent))
        elif c.tag == "a" and c.attrs.get("id") == "em":
            out.append(Run(email, href="mailto:" + email))
        elif c.tag == "a":
            href = urljoin(DESIGNED_URL, c.attrs.get("href") or "")
            out.extend(Run(r.text, href=href) for r in _runs(c, email))
        elif c.tag == "br":
            out.append(Run(br=True))
        else:
            raise Unknown(c)
    return out


def _clean(runs):
    """Trim the outer whitespace of a run list (it is layout, not words)."""
    runs = [r for r in runs if r.br or r.text]
    if runs and not runs[0].br:
        runs[0].text = runs[0].text.lstrip()
    if runs and not runs[-1].br:
        runs[-1].text = runs[-1].text.rstrip()
    return [r for r in runs if r.br or r.text]


def text_of(runs):
    return "".join("\n" if r.br else r.text for r in runs)


def model(source_html):
    """The printed résumé as plain data, for the Word files: name, title, contact and sections.
    Every printed word of resume/index.html must land in it, which the build checks."""
    src, sheet, email = read_source(source_html)
    m = {"name": None, "title": [], "contact": [], "sections": [], "email": email}
    for n in only_tags(sheet):
        if n.is_("header", "top"):
            for c in find(n, lambda x: x.tag in ("h1", "p")):
                if c.tag == "h1":
                    m["name"] = words(c)
                elif c.is_("p", "title"):
                    m["title"] = _clean(_runs(c, email))
                elif c.is_("p", "contact"):
                    m["contact"] = [_clean(_runs(item, email)) for item in only_tags(c)]
                else:
                    raise Unknown(c)
        elif n.tag == "h2":
            m["sections"].append({"heading": words(n), "blocks": []})
        elif "no-print" in n.classes:
            continue
        else:
            if not m["sections"]:
                raise Unknown(n)
            blocks = m["sections"][-1]["blocks"]
            if n.tag == "p" and n.classes <= {"summary"}:
                blocks.append(("para", _clean(_runs(n, email))))
            elif n.is_("p", "edu-lines"):
                lines, cur = [], []
                for r in _runs(n, email):
                    if r.br:
                        lines.append(_clean(cur))
                        cur = []
                    else:
                        cur.append(r)
                lines.append(_clean(cur))
                blocks.append(("lines", lines))
            elif n.is_("dl", "skills"):
                kids = only_tags(n)
                if len(kids) % 2 or any(k.tag != ("dt", "dd")[i % 2] for i, k in enumerate(kids)):
                    raise Unknown(n)
                blocks.append(("terms", [(_clean(_runs(kids[i], email)), _clean(_runs(kids[i + 1], email)))
                                         for i in range(0, len(kids), 2)]))
            elif n.is_("section", "role"):
                head = the_one(n, lambda x: x.is_("div", "role-head"), "role head")
                role = {"title": words(the_one(head, lambda x: x.tag == "h3", "role title")),
                        "when": words(the_one(head, lambda x: x.tag == "span", "role dates")),
                        "org": words(the_one(n, lambda x: x.is_("p", "org"), "role org")),
                        "bullets": [_clean(_runs(li, email)) for li in only_tags(the_one(n, lambda x: x.tag == "ul", "role list"))]}
                if [k.tag for k in only_tags(n)] != ["div", "p", "ul"]:
                    raise Unknown(n)
                blocks.append(("role", role))
            elif n.is_("div", "edu"):
                spans = only_tags(n)
                if [s.tag for s in spans] != ["span", "span"] or "edu-when" not in spans[1].classes:
                    raise Unknown(n)
                blocks.append(("school", {"name": words(spans[0]), "when": words(spans[1])}))
            elif n.tag == "ul":
                blocks.append(("list", [_clean(_runs(li, email)) for li in only_tags(n)], "certs" in n.classes))
            else:
                raise Unknown(n)
    return m


def model_words(m):
    parts = [m["name"], text_of(m["title"])] + [text_of(c) for c in m["contact"]]
    for s in m["sections"]:
        parts.append(s["heading"])
        for b in s["blocks"]:
            kind = b[0]
            if kind == "para":
                parts.append(text_of(b[1]))
            elif kind == "lines":
                parts.extend(text_of(line) for line in b[1])
            elif kind == "terms":
                for dt, dd in b[1]:
                    parts.extend([text_of(dt), text_of(dd)])
            elif kind == "role":
                r = b[1]
                parts.extend([r["title"], r["when"], r["org"]] + [text_of(x) for x in r["bullets"]])
            elif kind == "school":
                parts.extend([b[1]["name"], b[1]["when"]])
            elif kind == "list":
                parts.extend(text_of(x) for x in b[1])
    return " ".join(" ".join(parts).split())


def printed_words(source_html):
    """Every word a printout of resume/index.html shows, in order (both views print the same words)."""
    _, sheet, email = read_source(source_html)
    want = words(sheet, printed=True, email=email)
    same_words("the data for the Word files", want, model_words(model(source_html)))
    return want


# ------------------------------------------------------------------ Word

INK, DIM, BODY, PINK, VIOLET, BLUE, LINE, WASH, RULE = (
    "1B1328", "574A6B", "3A2F49", "D92F78", "8A4FE0", "2F6FE0", "EBE4F3", "FAF7FD", "B9CFF7")
PPR = ("pStyle keepNext keepLines pageBreakBefore framePr widowControl numPr suppressLineNumbers pBdr shd tabs "
       "suppressAutoHyphens kinsoku wordWrap overflowPunct topLinePunct autoSpaceDE autoSpaceDN bidi adjustRightInd "
       "snapToGrid spacing ind contextualSpacing mirrorIndents suppressOverlap jc textDirection textAlignment "
       "textboxTightWrap outlineLvl divId cnfStyle rPr sectPr pPrChange").split()
RPR = ("rStyle rFonts b bCs i iCs caps smallCaps strike dstrike outline shadow emboss imprint noProof snapToGrid "
       "vanish webHidden color spacing w kern position sz szCs highlight u effect bdr shd fitText vertAlign rtl cs em "
       "lang eastAsianLayout specVanish oMath").split()
TCPR = ("cnfStyle tcW gridSpan hMerge vMerge tcBorders shd noWrap tcMar textDirection tcFitText vAlign hideMark "
        "headers cellIns cellDel cellMerge tcPrChange").split()
TBLPR = ("tblStyle tblpPr tblOverlap bidiVisual tblStyleRowBandSize tblStyleColBandSize tblW jc tblCellSpacing "
         "tblInd tblBorders shd tblLayout tblCellMar tblLook tblCaption tblDescription tblPrChange").split()


def _gradient_png(width=1600, height=8):
    """The pink-violet-blue strip across the top of the designed sheet, as a PNG (no imaging library)."""
    stops = [(0.0, (0xff, 0x5c, 0xa8)), (0.48, (0xc4, 0x6b, 0xff)), (1.0, (0x4c, 0x8d, 0xff))]

    def at(t):
        for (t0, c0), (t1, c1) in zip(stops, stops[1:]):
            if t <= t1:
                k = (t - t0) / (t1 - t0)
                return bytes(int(round(a + (b - a) * k)) for a, b in zip(c0, c1))
        return bytes(stops[-1][1])

    row = b"".join(at(x / (width - 1)) for x in range(width))
    raw = b"".join(b"\x00" + row for _ in range(height))

    def chunk(kind, data):
        return struct.pack(">I", len(data)) + kind + data + struct.pack(">I", zlib.crc32(kind + data) & 0xffffffff)

    return (b"\x89PNG\r\n\x1a\n" + chunk(b"IHDR", struct.pack(">IIBBBBB", width, height, 8, 2, 0, 0, 0))
            + chunk(b"IDAT", zlib.compress(raw, 9)) + chunk(b"IEND", b""))


class Word(object):
    """A python-docx document plus the few raw-XML helpers python-docx doesn't have."""

    def __init__(self, margins, font, size):
        import docx
        from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
        from docx.enum.text import WD_TAB_ALIGNMENT, WD_LINE_SPACING, WD_ALIGN_PARAGRAPH
        from docx.opc.constants import RELATIONSHIP_TYPE
        from docx.oxml import OxmlElement
        from docx.oxml.ns import qn
        from docx.shared import Inches, Pt, RGBColor
        self.Inches, self.Pt, self.RGB, self.qn, self.OX = Inches, Pt, RGBColor, qn, OxmlElement
        self.TAB, self.LINE, self.ALIGN, self.RT = WD_TAB_ALIGNMENT, WD_LINE_SPACING, WD_ALIGN_PARAGRAPH, RELATIONSHIP_TYPE
        self.TABLE_ALIGN, self.VALIGN = WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
        self.d = docx.Document()
        sec = self.d.sections[0]
        sec.page_width, sec.page_height = Inches(8.5), Inches(11)
        top, right, bottom, left = margins
        sec.top_margin, sec.right_margin, sec.bottom_margin, sec.left_margin = (
            Inches(top), Inches(right), Inches(bottom), Inches(left))
        self.width = 8.5 - left - right
        normal = self.d.styles["Normal"]
        self.plain_font(normal.element.get_or_add_rPr(), font)
        normal.font.size = Pt(size)
        normal.paragraph_format.space_before = Pt(0)
        normal.paragraph_format.space_after = Pt(0)
        self.font = font

    # raw XML
    def el(self, tag, **attrs):
        e = self.OX(tag)
        for k, v in attrs.items():
            e.set(self.qn("w:" + k), str(v))
        return e

    def put(self, parent, child, order):
        name = child.tag.split("}")[1]
        for old in parent.findall(child.tag):
            parent.remove(old)
        rank = order.index(name)
        for e in parent:
            other = e.tag.split("}")[1]
            if other in order and order.index(other) > rank:
                e.addprevious(child)
                return child
        parent.append(child)
        return child

    def plain_font(self, rpr, font):
        """Name the font outright: drop the theme fonts the default template would use instead."""
        fonts = self.el("w:rFonts", ascii=font, hAnsi=font, cs=font, eastAsia=font)
        self.put(rpr, fonts, RPR)

    # text
    def para(self, box, before=0, after=0, line=None, keep=False, align=None, exact=None):
        p = box.add_paragraph()
        pf = p.paragraph_format
        pf.space_before, pf.space_after = self.Pt(before), self.Pt(after)
        if line:
            pf.line_spacing = line
        if exact:
            pf.line_spacing_rule = self.LINE.EXACTLY
            pf.line_spacing = self.Pt(exact)
        if keep:
            pf.keep_with_next = True
        if align:
            pf.alignment = getattr(self.ALIGN, align)
        return p

    def run(self, p, text, size, color=None, bold=False, caps=False, spacing=None, font=None):
        r = p.add_run(text)
        r.font.size = self.Pt(size)
        if bold:
            r.font.bold = True
        if color:
            r.font.color.rgb = self.RGB.from_string(color)
        if caps:
            r.font.all_caps = True
        if font:
            self.plain_font(r._r.get_or_add_rPr(), font)
        if spacing:
            self.put(r._r.get_or_add_rPr(), self.el("w:spacing", val=int(round(spacing * 20))), RPR)
        return r

    def link(self, p, text, url, size, color, bold=False):
        rid = p.part.relate_to(url, self.RT.HYPERLINK, is_external=True)
        h = self.OX("w:hyperlink")
        h.set(self.qn("r:id"), rid)
        r = self.OX("w:r")
        rpr = self.OX("w:rPr")
        if bold:
            rpr.append(self.el("w:b"))
        rpr.append(self.el("w:color", val=color))
        rpr.append(self.el("w:sz", val=int(round(size * 2))))
        rpr.append(self.el("w:u", val="single"))
        r.append(rpr)
        t = self.OX("w:t")
        t.text = text
        t.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
        r.append(t)
        h.append(r)
        p._p.append(h)

    def runs(self, p, runs, size, color=None, bold=False, accent=None, link_color=None):
        for r in runs:
            if r.br:
                p.add_run().add_break()
            elif r.href:
                self.link(p, r.text, r.href, size, link_color or color or "000000", bold=bold)
            else:
                self.run(p, r.text, size, accent if (r.accent and accent) else color, bold=bold)

    def border_below(self, p_or_style_ppr, color, size_eighths, space=1):
        bdr = self.OX("w:pBdr")
        bdr.append(self.el("w:bottom", val="single", sz=size_eighths, space=space, color=color))
        self.put(p_or_style_ppr, bdr, PPR)

    def right_tab(self, p, inches):
        p.paragraph_format.tab_stops.add_tab_stop(self.Inches(inches), self.TAB.RIGHT)

    def bullets(self, color, left=0.2, hang=0.14):
        """A bullet list definition (•) in the given color; returns its numId."""
        numbering = self.d.part.numbering_part.element
        q = self.qn
        aid = max([int(a.get(q("w:abstractNumId"))) for a in numbering.findall(q("w:abstractNum"))] + [0]) + 1
        nid = max([int(n.get(q("w:numId"))) for n in numbering.findall(q("w:num"))] + [0]) + 1
        a = self.OX("w:abstractNum")
        a.set(q("w:abstractNumId"), str(aid))
        a.append(self.el("w:multiLevelType", val="singleLevel"))
        lvl = self.OX("w:lvl")
        lvl.set(q("w:ilvl"), "0")
        for tag, val in (("w:start", 1), ("w:numFmt", "bullet"), ("w:lvlText", "•"), ("w:lvlJc", "left")):
            lvl.append(self.el(tag, val=val))
        ppr = self.OX("w:pPr")
        ppr.append(self.el("w:ind", left=int(left * 1440), hanging=int(hang * 1440)))
        lvl.append(ppr)
        rpr = self.OX("w:rPr")
        rpr.append(self.el("w:rFonts", ascii=self.font, hAnsi=self.font, cs=self.font))
        rpr.append(self.el("w:color", val=color))
        lvl.append(rpr)
        a.append(lvl)
        nums = numbering.findall(q("w:num"))
        if nums:
            nums[0].addprevious(a)
        else:
            numbering.append(a)
        n = self.OX("w:num")
        n.set(q("w:numId"), str(nid))
        n.append(self.el("w:abstractNumId", val=aid))
        numbering.append(n)
        return nid

    def bullet(self, box, numid, runs, size, color, after, line=None, accent=None, link_color=None):
        p = self.para(box, after=after, line=line)
        num = self.OX("w:numPr")
        num.append(self.el("w:ilvl", val=0))
        num.append(self.el("w:numId", val=numid))
        self.put(p._p.get_or_add_pPr(), num, PPR)
        self.runs(p, runs, size, color, accent=accent, link_color=link_color)
        return p

    # tables
    def table(self, box, widths, rows=1, margins=(0, 0, 0, 0), spacing=None):
        t = box.add_table(rows=rows, cols=len(widths))
        t.autofit = False
        t.alignment = self.TABLE_ALIGN.LEFT
        tblpr = t._tbl.tblPr
        if spacing is not None:
            self.put(tblpr, self.el("w:tblCellSpacing", w=int(spacing * 1440), type="dxa"), TBLPR)
        self.put(tblpr, self.el("w:tblW", w=int(round(sum(widths) * 1440)), type="dxa"), TBLPR)
        self.put(tblpr, self.el("w:tblInd", w=0, type="dxa"), TBLPR)
        borders = self.OX("w:tblBorders")
        for side in ("top", "left", "bottom", "right", "insideH", "insideV"):
            borders.append(self.el("w:" + side, val="nil"))
        self.put(tblpr, borders, TBLPR)
        self.put(tblpr, self.el("w:tblLayout", type="fixed"), TBLPR)
        mar = self.OX("w:tblCellMar")
        for side, v in zip(("top", "left", "bottom", "right"), margins):
            mar.append(self.el("w:" + side, w=int(v * 1440), type="dxa"))
        self.put(tblpr, mar, TBLPR)
        for col, w in zip(t._tbl.tblGrid.findall(self.qn("w:gridCol")), widths):
            col.set(self.qn("w:w"), str(int(round(w * 1440))))
        for row in t.rows:
            for cell, w in zip(row.cells, widths):
                cell.width = self.Inches(w)
        return t

    def box(self, cell, fill, border):
        tcpr = cell._tc.get_or_add_tcPr()
        b = self.OX("w:tcBorders")
        for side in ("top", "left", "bottom", "right"):
            b.append(self.el("w:" + side, val="single", sz=4, space=0, color=border))
        self.put(tcpr, b, TCPR)
        self.put(tcpr, self.el("w:shd", val="clear", color="auto", fill=fill), TCPR)

    def gap(self, points):
        """Word merges two tables that touch, so a thin empty paragraph keeps them apart."""
        return self.para(self.d, exact=points)

    def picture(self, p, data, width, height=None, alt=""):
        r = p.add_run()
        stream = io.BytesIO(data)
        r.add_picture(stream, width=width, height=height) if height else r.add_picture(stream, width=width)
        for pr in r._r.iter("{http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing}docPr"):
            pr.set("descr", alt)
            pr.set("name", alt or "decoration")
        return r

    def finish(self, title):
        from docx.opc.constants import RELATIONSHIP_TYPE as RT
        cp = self.d.core_properties
        now = datetime.now(timezone.utc).replace(microsecond=0, tzinfo=None)
        cp.title, cp.subject, cp.author, cp.last_modified_by = title, "Résumé", "Michael Purpura", "Michael Purpura"
        cp.comments = "Built from https://okay-cool.net/resume/ by tools/resume_build.py"
        cp.keywords, cp.category, cp.revision = "", "", 1
        cp.created = cp.modified = now
        rels = self.d.part.package.rels
        store = getattr(rels, "_rels", rels)
        for rid, rel in list(store.items()):
            if rel.reltype == RT.THUMBNAIL:  # the default template's blank-page thumbnail
                store.pop(rid)
        zoom = self.d.settings.element.find(self.qn("w:zoom"))
        if zoom is not None and zoom.get(self.qn("w:percent")) is None:
            zoom.set(self.qn("w:percent"), "100")  # the template leaves out this required attribute
        out = io.BytesIO()
        self.d.save(out)
        return out.getvalue()


def designed_docx(m, root):
    """Word, dressed like the designed sheet: gradient strip, serif name, QR code, boxed roles, chips."""
    w = Word((0.42, 0.5, 0.4, 0.5), "Calibri", 9.6)
    d, pt, inch = w.d, w.Pt, w.Inches
    top = d.paragraphs[0] if d.paragraphs else w.para(d)
    w.picture(top, _gradient_png(), inch(w.width), pt(4))
    top.paragraph_format.space_after = pt(9)
    head = w.table(d, [w.width - 1.15, 1.15])
    left, right = head.rows[0].cells
    p = left.paragraphs[0]
    w.run(p, m["name"], 30, INK, font="Georgia")
    p.paragraph_format.space_after = pt(3)
    p = w.para(left, after=4)
    w.runs(p, m["title"], 10.4, INK, bold=True, accent=PINK)
    p = w.para(left)
    for i, item in enumerate(m["contact"]):
        if i:
            w.run(p, " ", 9.2)
        w.run(p, "◆ ", 7, VIOLET)
        w.runs(p, item, 9.2, DIM, link_color=INK)
    right.vertical_alignment = w.VALIGN.CENTER
    p = right.paragraphs[0]
    p.paragraph_format.alignment = w.ALIGN.RIGHT
    with open(os.path.join(root, QR), "rb") as f:
        w.picture(p, f.read(), inch(0.95), alt="QR code that opens okay-cool.net")
    rule = w.para(d, after=2, exact=4)
    w.border_below(rule._p.get_or_add_pPr(), LINE, 6)
    dots = w.bullets(PINK)
    inner = w.width - 0.28
    for s in m["sections"]:
        h = w.para(d, before=10, after=5, keep=True)
        w.run(h, s["heading"], 8.6, BLUE, bold=True, caps=True, spacing=1.4)
        w.border_below(h._p.get_or_add_pPr(), RULE, 6, space=2)
        for b in s["blocks"]:
            kind = b[0]
            if kind == "para":
                w.runs(w.para(d, after=3, line=1.15), b[1], 9.8, INK)
            elif kind == "terms":
                t = w.table(d, [1.6, w.width - 1.6], rows=len(b[1]), margins=(0.012, 0, 0.012, 0.06))
                for row, (dt, dd) in zip(t.rows, b[1]):
                    w.runs(row.cells[0].paragraphs[0], dt, 9.2, VIOLET, bold=True)
                    w.runs(row.cells[1].paragraphs[0], dd, 9.2, BODY)
            elif kind == "role":
                r = b[1]
                t = w.table(d, [w.width], margins=(0.1, 0.14, 0.05, 0.14))
                cell = t.rows[0].cells[0]
                w.box(cell, WASH, LINE)
                p = cell.paragraphs[0]
                w.right_tab(p, inner)
                w.run(p, r["title"], 10.2, INK, bold=True)
                w.run(p, "\t", 10.2)
                w.run(p, r["when"], 8.6, PINK, bold=True)
                w.run(w.para(cell, after=3), r["org"], 9.2, DIM)
                for x in r["bullets"]:
                    w.bullet(cell, dots, x, 9.4, INK, after=2.4, line=1.1)
                w.gap(7)
            elif kind == "school":
                p = w.para(d, after=1)
                w.right_tab(p, w.width)
                w.run(p, b[1]["name"], 9.6, INK, bold=True)
                w.run(p, "\t", 9.6)
                w.run(p, b[1]["when"], 8.6, PINK, bold=True)
            elif kind == "lines":
                p = w.para(d, after=2)
                for i, line in enumerate(b[1]):
                    if i:
                        p.add_run().add_break()
                    w.runs(p, line, 9.4, BODY)
            elif kind == "list" and b[2]:
                chips = [text_of(x) for x in b[1]]
                widths = [0.067 * len(c) + 0.26 for c in chips]
                t = w.table(d, widths, margins=(0.03, 0.1, 0.03, 0.1), spacing=0.04)
                for cell, x in zip(t.rows[0].cells, b[1]):
                    w.box(cell, WASH, LINE)
                    w.runs(cell.paragraphs[0], x, 9, INK, bold=True)
            elif kind == "list":
                for x in b[1]:
                    w.bullet(d, dots, x, 9.4, INK, after=2.4)
    return w.finish(m["name"] + " · Résumé")


def plain_docx(m):
    """Word, plain: one column, Calibri, built-in heading styles and real bullet lists. No tables,
    no pictures, so screen readers and applicant-tracking systems read it in order."""
    w = Word((0.55, 0.6, 0.55, 0.6), "Calibri", 10.5)
    d, pt = w.d, w.Pt
    for name, size, before, after in (("Heading 1", 12.5, 12, 4), ("Heading 2", 11, 9, 0)):
        st = d.styles[name]
        rpr = st.element.get_or_add_rPr()
        w.plain_font(rpr, "Calibri")
        st.font.size, st.font.bold, st.font.italic = pt(size), True, False
        st.font.color.rgb = w.RGB.from_string("000000")
        for c in rpr.findall(w.qn("w:color")):
            for k in ("themeColor", "themeShade", "themeTint"):
                c.attrib.pop(w.qn("w:" + k), None)
        st.paragraph_format.space_before, st.paragraph_format.space_after = pt(before), pt(after)
        st.paragraph_format.keep_with_next = True
        if name == "Heading 1":
            w.border_below(st.element.get_or_add_pPr(), "000000", 8, space=2)
    dots = w.bullets("000000", left=0.25, hang=0.17)
    first = d.paragraphs[0] if d.paragraphs else w.para(d)
    w.run(first, m["name"], 20, "000000", bold=True)
    first.paragraph_format.space_after = pt(2)
    w.runs(w.para(d, after=3), m["title"], 10.5, "000000", bold=True)
    p = w.para(d, after=2)
    for i, item in enumerate(m["contact"]):
        if i:
            w.run(p, "  ", 10.5)
        w.runs(p, item, 10.5, "000000", link_color="0B4FB3")
    for s in m["sections"]:
        d.add_paragraph(s["heading"], style="Heading 1")
        for b in s["blocks"]:
            kind = b[0]
            if kind == "para":
                w.runs(w.para(d, after=4, line=1.1), b[1], 10.5, "000000")
            elif kind == "terms":
                for dt, dd in b[1]:
                    w.runs(w.para(d, keep=True), dt, 10.5, "000000", bold=True)
                    w.runs(w.para(d, after=5), dd, 10.5, "000000")
            elif kind == "role":
                r = b[1]
                d.add_paragraph(r["title"], style="Heading 2")
                w.run(w.para(d, keep=True), r["when"], 10.5, "000000")
                w.run(w.para(d, after=3, keep=True), r["org"], 10.5, "000000")
                for x in r["bullets"]:
                    w.bullet(d, dots, x, 10.5, "000000", after=2, line=1.05)
            elif kind == "school":
                d.add_paragraph(b[1]["name"], style="Heading 2")
                w.run(w.para(d, keep=True), b[1]["when"], 10.5, "000000")
            elif kind == "lines":
                p = w.para(d, before=3, after=2)
                for i, line in enumerate(b[1]):
                    if i:
                        p.add_run().add_break()
                    w.runs(p, line, 10.5, "000000")
            elif kind == "list":
                for x in b[1]:
                    w.bullet(d, dots, x, 10.5, "000000", after=2)
    return w.finish(m["name"] + " · Résumé (plain)")


def docx_words(data):
    import docx
    from docx.oxml.ns import qn
    from docx.text.paragraph import Paragraph
    doc = docx.Document(io.BytesIO(data))
    texts = []

    def walk(el):
        for child in el.iterchildren():
            if child.tag == qn("w:p"):
                texts.append(Paragraph(child, None).text)
            elif child.tag == qn("w:tbl"):
                for row in child.iterchildren(qn("w:tr")):
                    for cell in row.iterchildren(qn("w:tc")):
                        walk(cell)

    walk(doc.element.body)
    return " ".join(" ".join(texts).replace("◆", " ").split())


# ------------------------------------------------------------------ PDF, printed by a headless browser

class _Quiet(http.server.SimpleHTTPRequestHandler):
    extensions_map = dict(http.server.SimpleHTTPRequestHandler.extensions_map,
                          **{".js": "text/javascript", ".woff2": "font/woff2", ".svg": "image/svg+xml"})

    def log_message(self, *a):
        pass


class _Tab(object):
    def __init__(self, port):
        import websocket
        req = urllib.request.Request("http://127.0.0.1:%d/json/new?about:blank" % port, method="PUT")
        info = json.load(urllib.request.urlopen(req, timeout=10))
        self.ws = websocket.create_connection(info["webSocketDebuggerUrl"], suppress_origin=True, timeout=90)
        self.n = 0

    def call(self, method, **params):
        self.n += 1
        self.ws.send(json.dumps({"id": self.n, "method": method, "params": params}))
        while True:
            msg = json.loads(self.ws.recv())
            if msg.get("id") == self.n:
                if "error" in msg:
                    raise RuntimeError("%s: %s" % (method, msg["error"]))
                return msg.get("result", {})

    def js(self, expr, wait=False):
        r = self.call("Runtime.evaluate", expression=expr, returnByValue=True, awaitPromise=wait)
        if "exceptionDetails" in r:
            raise RuntimeError("page script failed: %s" % r["exceptionDetails"])
        return r.get("result", {}).get("value")


def find_browser(explicit=None):
    for b in ([explicit] if explicit else []) + BROWSERS:
        if b and (os.path.isfile(b) or shutil.which(b)):
            return b if os.path.isfile(b) else shutil.which(b)
    raise SystemExit("no headless browser found to print the PDFs; pass --browser, or --no-pdf to skip them")


def print_pdfs(root, browser):
    """Print the designed and plain views of resume/index.html. Returns {path: pdf bytes}."""
    srv = http.server.ThreadingHTTPServer(("127.0.0.1", 0), functools.partial(_Quiet, directory=root))
    threading.Thread(target=srv.serve_forever, daemon=True).start()
    site = "http://127.0.0.1:%d" % srv.server_address[1]
    profile = tempfile.mkdtemp(prefix="resume-pdf-")
    proc = subprocess.Popen([browser, "--headless=new", "--remote-debugging-port=0", "--user-data-dir=" + profile,
                             "--no-first-run", "--no-default-browser-check", "--disable-gpu", "--disable-extensions",
                             "--no-proxy-server", "about:blank"], stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
    out = {}
    try:
        port_file, t0 = os.path.join(profile, "DevToolsActivePort"), time.time()
        while not (os.path.exists(port_file) and open(port_file).read().strip()):
            if time.time() - t0 > 30 or proc.poll() is not None:
                raise RuntimeError("the browser didn't start: " + browser)
            time.sleep(0.2)
        port = int(open(port_file).read().split()[0])
        for path, view in ((DESIGNED_PDF, ""), (PLAIN_PDF, "#plain")):
            tab = _Tab(port)
            try:
                tab.call("Emulation.setEmulatedMedia", features=[{"name": "prefers-color-scheme", "value": "light"}])
                tab.call("Page.enable")
                tab.call("Page.navigate", url=site + "/resume/" + view)
                t0 = time.time()
                while tab.js("document.readyState") != "complete":
                    if time.time() - t0 > 30:
                        raise RuntimeError("the résumé page didn't finish loading")
                    time.sleep(0.1)
                tab.js("document.fonts.ready.then(() => true)", wait=True)
                time.sleep(0.3)
                state = tab.js("({plain: document.getElementById('view-plain').checked,"
                               " email: document.getElementById('em').textContent})")
                if state["plain"] != bool(view) or "@" not in state["email"]:
                    raise RuntimeError("the page isn't ready to print the %s view: %r" % (view or "designed", state))
                params = dict(preferCSSPageSize=True, printBackground=True, generateTaggedPDF=True,
                              generateDocumentOutline=True)
                try:
                    data = tab.call("Page.printToPDF", **params)["data"]
                except RuntimeError:
                    for k in ("generateTaggedPDF", "generateDocumentOutline"):
                        params.pop(k)
                    data = tab.call("Page.printToPDF", **params)["data"]
                out[path] = base64.b64decode(data)
            finally:
                tab.ws.close()
    finally:
        proc.terminate()
        try:
            proc.wait(10)
        except subprocess.TimeoutExpired:
            proc.kill()
        shutil.rmtree(profile, ignore_errors=True)
        srv.shutdown()
    return out


def pdf_facts(data):
    import pymupdf
    doc = pymupdf.open(stream=data, filetype="pdf")
    text = "".join(page.get_text() for page in doc)
    return doc.page_count, text


def squash(text):
    """Compare PDFs by their characters: a printout re-flows lines and the designed view spaces out
    its headings, so whitespace and case can't be trusted, but every character must be there."""
    return re.sub(r"\s+", "", text.replace("◆", "")).casefold()


def label_pdf(data, title):
    """Name the PDF's author and title (the browser only fills in its own name)."""
    import pymupdf
    doc = pymupdf.open(stream=data, filetype="pdf")
    meta = dict(doc.metadata or {})
    meta.update(title=title, author="Michael Purpura", subject="Résumé", creator="okay-cool.net")
    doc.set_metadata(meta)
    return doc.tobytes(garbage=0, deflate=True)


# ------------------------------------------------------------------ check and build

def check_files(root, want, files):
    """Read each built file back. Returns a list of problems (empty when everything matches)."""
    problems = []
    flat = squash(want)
    for path, data in sorted(files.items()):
        if data is None:
            problems.append("%s is missing" % path)
        elif path.endswith(".pdf"):
            pages, text = pdf_facts(data)
            if pages != PDF_PAGES:
                problems.append("%s has %d pages, not %d" % (path, pages, PDF_PAGES))
            if squash(text) != flat:
                sm = difflib.SequenceMatcher(None, flat, squash(text), autojunk=False)
                bad = ["%r -> %r" % (flat[a:b][:60], squash(text)[c:e][:60])
                       for op, a, b, c, e in sm.get_opcodes() if op != "equal"][:4]
                problems.append("%s doesn't carry the résumé's words: %s" % (path, "; ".join(bad)))
        elif path.endswith(".docx"):
            got = docx_words(data)
            if got != want:
                diff = list(difflib.unified_diff(want.split(), got.split(), "resume/index.html", path, lineterm="", n=2))
                problems.append("%s doesn't match word for word:\n%s" % (path, "\n".join(diff[:30])))
    return problems


def main():
    ap = argparse.ArgumentParser(description="Build the plain page and the four résumé downloads from resume/index.html.")
    ap.add_argument("--check", action="store_true", help="change nothing; exit 1 if any file is out of date")
    ap.add_argument("--root", default=REPO, help="the site checkout (default: the one this tool sits in)")
    ap.add_argument("--browser", help="headless Chromium that prints the PDFs (default: Brave, then Chrome)")
    ap.add_argument("--no-pdf", action="store_true", help="leave the PDFs alone (no browser here)")
    args = ap.parse_args()
    root = os.path.abspath(args.root)
    with open(os.path.join(root, SOURCE), encoding="utf-8") as f:
        source = f.read()
    want = printed_words(source)
    page = plain_page(source)
    m = model(source)

    def on_disk(path):
        full = os.path.join(root, path)
        return open(full, "rb").read() if os.path.exists(full) else None

    if args.check:
        stale = [] if on_disk(PLAIN_PAGE) == page.encode("utf-8") else ["%s is out of date" % PLAIN_PAGE]
        stale += check_files(root, want, {p: on_disk(p) for p in (DESIGNED_DOCX, PLAIN_DOCX)
                                          + (() if args.no_pdf else (DESIGNED_PDF, PLAIN_PDF))})
        print("\n".join(stale) if stale else "every résumé file matches resume/index.html")
        return 1 if stale else 0

    built = {DESIGNED_DOCX: designed_docx(m, root), PLAIN_DOCX: plain_docx(m)}
    if not args.no_pdf:
        printed = print_pdfs(root, find_browser(args.browser))
        built[DESIGNED_PDF] = label_pdf(printed[DESIGNED_PDF], m["name"] + " · Résumé")
        built[PLAIN_PDF] = label_pdf(printed[PLAIN_PDF], m["name"] + " · Résumé (plain)")
    problems = check_files(root, want, built)
    if problems:
        print("Nothing written:\n" + "\n".join(problems), file=sys.stderr)
        return 1
    os.makedirs(os.path.join(root, os.path.dirname(PLAIN_PAGE)), exist_ok=True)
    with open(os.path.join(root, PLAIN_PAGE), "w", encoding="utf-8") as f:
        f.write(page)
    for path, data in built.items():
        with open(os.path.join(root, path), "wb") as f:
            f.write(data)
    print("%s: the page's %d words. %s: the %d printed words (the email spelled out, no screen-only footer),"
          " in order; PDFs %d pages each." % (PLAIN_PAGE, len(words(read_source(source)[1]).split()),
                                             ", ".join(sorted(built)), len(want.split()), PDF_PAGES))
    return 0


if __name__ == "__main__":
    sys.exit(main())
